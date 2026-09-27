"""Bounded offline preparation of pinned September 2026 field sources.

No runtime admission. All original records receive dispositions, including
metadata, formulas, missingness and held conflicts. Outputs require review.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import re
import zipfile
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from agronomy_agent.field_data import MAX_CELL_CHARS, MAX_COLUMNS, MAX_ROWS, MAX_SOURCE_BYTES, MAX_ZIP_EXPANDED_BYTES, MAX_ZIP_MEMBERS, parse_table, safe_number

VERSION = "field-source-preparation-v1"
RAW_DIRECTORY = "data/raw/field_data_research/20260927-expansion"
# Copied from frozen acquisition-receipts.json; never learned from supplied bytes.
PINS = {
    "ca-borealis-bean-ayd-field/AYD_Field_Data_4environments_YRKPP.tab": "7ab51a52725515fe9de93166fb4146aa5b09dd07ef6e3b96a40b8c411767553d",
    "ca-borealis-bean-ayd-field/AYD_AM_Codebook_YRKPP.docx": "3dda0a8c7edeb15a59f781e6af541d9535b016e11477eee32c99242ec306bd09",
    "ca-borealis-onion-tissue-yield/100A_README.txt": "492d9b7cbf7723f447251af329215d23c8ad68a10f1512e68a935fa57234eb64",
    "ca-borealis-onion-tissue-yield/growerFieldOnion_2024And2025.tab": "4474725844580681017893421779038ba0080f3313dccb863397eb12fae0878a",
    "ca-borealis-onion-tissue-yield/micronutrientTrialOnion_2024And2025.tab": "e0f3057bc162a59c88c1c6b831894cb72adb23093f7957b8dc0b457eb1bc121e",
    "ca-borealis-onion-tissue-yield/stemphyliumLeafBlightOnionYield_2024And2025.tab": "324f50eb37323495c770241efcc42eb1318a4c27e3217c2b1e635e3fd1181bb1",
    "us-arkansas-soybean-p-39/Slaton et al., 2021 Dataset.xlsx": "64a9be4054f87b7e847483915b9326cf189a721a4cf72bcbbb9114f9c561a852",
}
NUTRIENTS = ["Nitrogen (%)", "Phosphorus (%)", "Potassium (%)", "Magnesium (%)", "Calcium (%)", "Zinc (ppm)", "Manganese (ppm)", "Copper (ppm)", "Iron (ppm)", "Boron (ppm)", "Sulfur (%)"]
ONION = "ca-borealis-onion-tissue-yield/"
PROFILES = [
    dict(id="bean", path="ca-borealis-bean-ayd-field/AYD_Field_Data_4environments_YRKPP.tab", sheet="Sheet1", group=["Location"], context=["Year", "Location", "Plot", "Block", "iBLOCK", "Entry", "Cultivar"], measurements={"YD": "kg/ha", "SW": "g/100 seeds", "DF": "days", "DM": "days", "PH": "cm", "HR": "score 1-5"}, url="https://doi.org/10.5683/SP3/FD81LR", license="CC-BY-4.0", dictionary="ca-borealis-bean-ayd-field/AYD_AM_Codebook_YRKPP.docx", notes="Six measured traits only. YD and SW source-adjusted to 18% moisture. RP, YGD, SGR, YDH, SN, YDHR derived and excluded. Environments are not independent farms. Codebook Elora 2016 end-date typo unresolved; no dates inferred."),
    dict(id="onion-grower", path=ONION+"growerFieldOnion_2024And2025.tab", sheet=None, group=["Field", "Year", "Method"], context=["Year", "Date", "Field", "Sampling round", "Method", "Scan point"], measurements={n: "%" if "(%)" in n else "ppm" for n in NUTRIENTS}, url="https://doi.org/10.5683/SP3/YJTMRB", license="CC-BY-4.0", dictionary=ONION+"100A_README.txt", notes="SGS laboratory values are observations; Picketa device estimates conservatively model_output. Methods sample the same leaves. Date windows retained as text; README windows disagree with some source strings and are not substituted."),
    dict(id="onion-trial-tissue", path=ONION+"micronutrientTrialOnion_2024And2025.tab", sheet=None, group=["Year", "Method"], context=["Date", "Year", "Scan_time", "Method", "Preplant", "Foliar", "Rep"], measurements={n: "%" if "(%)" in n else "ppm" for n in NUTRIENTS}, url="https://doi.org/10.5683/SP3/YJTMRB", license="CC-BY-4.0", dictionary=ONION+"100A_README.txt", notes="Bradford trial; year/method partitions are not farms. Literal schema differs from README (18 versus 22 columns; no Trt_ID). No invented plot joins. Date/year conflicts held. Partial dates remain text."),
    dict(id="onion-trial-yield", path=ONION+"stemphyliumLeafBlightOnionYield_2024And2025.tab", sheet=None, group=["Year"], context=["Year", "Preplant", "Foliar", "Rep"], measurements={"Jumbo bulb weight (kg)": "kg", "Jumbo bulb number": "count", "large bulb weight (kg)": "kg", "large bulb number": "count", "Medium bulb weight (kg)": "kg", "Medium bulb number": "count", "Small bulb weight (kg)": "kg", "Small bulb number": "count"}, url="https://doi.org/10.5683/SP3/YJTMRB", license="CC-BY-4.0", dictionary=ONION+"100A_README.txt", notes="Literal weights/counts from graded 2.32 m plot segment. Summed totals, scaled t/ha yields, percentages and derived disease indices excluded. Harvest date/geometry unknown; no per-area inference."),
    dict(id="soybean", path="us-arkansas-soybean-p-39/Slaton et al., 2021 Dataset.xlsx", sheet="raw_data", group=["Site-ID2"], context=["Index", "Trial Harvest Year (YYYY)", "Site ID", "Site-ID2", "Plot", "Rep", "Crop Cultivar", "Treatment P Fertilizer Rate (kg P/ha)", "Treatment Fertilizer Type", "Plant Sample Collection Date (MM/DD/YYYY)"], measurements={"Actual Yield (kg/ha)": "kg/ha", "Crop Moisture Content %": "%", "P Concentration (g/kg)": "g/kg", "K Concentration (g/kg)": "g/kg", "Ca Concentration (g/kg)": "g/kg", "Mg Concentration (g/kg)": "g/kg", "S Concentration (g/kg)": "g/kg"}, derived=["Actual Yield (kg/ha)"], url="https://agdatacommons.nal.usda.gov/articles/dataset/24667830", license="U.S. Public Domain", dictionary="us-arkansas-soybean-p-39/Slaton et al., 2021 Dataset.xlsx#dictionary", notes="Site-ID2 is trial identity, not farm independence. Dictionary describes Actual Yield as average: conservatively derived/interpretation. Tissue concentrations are observations; fraction/stage formulas excluded. SoyMAP predicted date excluded model_output. Soil replicate averages excluded; micronutrient dictionary/header unit conflicts avoided. No geometry imported."),
]


def _json(value: Any) -> bytes:
    return (json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"), allow_nan=False)+"\n").encode()


def _sha(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _text(value: Any) -> str:
    return "" if value is None else value.isoformat() if isinstance(value, (datetime, date)) else str(value)


def _archive(content: bytes) -> None:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        if len(members) > MAX_ZIP_MEMBERS or sum(m.file_size for m in members) > MAX_ZIP_EXPANDED_BYTES:
            raise ValueError("archive exceeds bounds")
        if any(m.flag_bits & 1 or (m.compress_size and m.file_size > m.compress_size * 200) for m in members):
            raise ValueError("archive encryption or expansion exceeds bounds")


def _read_table(content: bytes, sheet: str | None) -> tuple[list[str], list[dict], list[dict]]:
    metadata, records = [], []
    if sheet is None:
        reader = csv.reader(io.StringIO(content.decode("utf-8-sig"), newline=""), strict=True)
        header = next(reader)
        metadata.append({"sheet": "CSV", "record": 1, "disposition": "header"})
        for i, values in enumerate(reader, 2):
            if i > MAX_ROWS + 1 or len(values) != len(header):
                raise ValueError("row bounds or width mismatch")
            records.append({"sheet": "CSV", "record": i, "physical_line_end": reader.line_num, "values": values, "formulas": []})
    else:
        import openpyxl
        _archive(content)
        book = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            if sheet not in book.sheetnames:
                raise ValueError("required source sheet missing")
            total = 0
            for tab in book:
                if tab.max_row > MAX_ROWS + 1 or tab.max_column > MAX_COLUMNS:
                    raise ValueError("worksheet bounds exceeded")
                for i, cells in enumerate(tab.iter_rows(), 1):
                    total += 1
                    if total > MAX_ROWS + 1:
                        raise ValueError("workbook row bounds exceeded")
                    values = [_text(c.value) for c in cells]
                    if i == 1 and tab.title == sheet:
                        while values and not values[-1]:
                            values.pop()
                        header = values
                    if tab.title != sheet or i == 1:
                        metadata.append({"sheet": tab.title, "record": i, "disposition": "header" if i == 1 else "metadata_excluded"})
                        continue
                    if any(values[len(header):]):
                        raise ValueError("nonempty cells outside header")
                    records.append({"sheet": tab.title, "record": i, "values": values[:len(header)], "formulas": [j for j, c in enumerate(cells[:len(header)]) if c.data_type == "f"]})
        finally:
            book.close()
    if not header or len(header) > MAX_COLUMNS or len(set(header)) != len(header):
        raise ValueError("invalid header")
    if any(len(v) > MAX_CELL_CHARS for r in records for v in r["values"]):
        raise ValueError("cell exceeds bounds")
    return header, records, metadata


def _alias(name: str) -> str:
    return "source_" + re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")


def _prepare_profile(profile: dict, content: bytes) -> tuple[dict, dict[str, list[dict]]]:
    header, records, dispositions = _read_table(content, profile["sheet"])
    selected = profile["context"] + list(profile["measurements"])
    if not set(selected + profile["group"]) <= set(header):
        raise ValueError("required source columns missing")
    if len({_alias(n) for n in selected}) != len(selected):
        raise ValueError("ambiguous output aliases")
    raw_sha = _sha(content)
    groups = defaultdict(list)
    formula_count = Counter()
    for rec in records:
        values = dict(zip(header, rec["values"], strict=True))
        decision = {**{k: rec[k] for k in ("sheet", "record", "physical_line_end") if k in rec}, "disposition": "prepared", "null_cells": [], "formula_columns_excluded": []}
        formulas = {header[i] for i in rec["formulas"]}
        for name in sorted(formulas):
            formula_count[name] += 1
            decision["formula_columns_excluded"].append(name)
        if not any(values.values()):
            decision["disposition"] = "blank_excluded"
        elif any(not values[k].strip() or k in formulas for k in profile["group"]):
            decision.update(disposition="held", reason="missing_or_formula_study_identity")
        elif "Date" in values and re.match(r"^\d{4}-", values["Date"]) and values.get("Year") != values["Date"][:4]:
            decision.update(disposition="held", reason="date_year_conflict", source_date=values["Date"], source_year=values.get("Year"))
        if decision["disposition"] in {"blank_excluded", "held"}:
            dispositions.append(decision)
            continue
        group = " | ".join(values[n] for n in profile["group"])
        row = {"study_unit": group, "source_row": f"sha256:{raw_sha}#{rec['sheet']}!record={rec['record']}"}
        evidence = "model_output" if values.get("Method") == "Picketa_LENS" else "observation"
        if "Method" in values and values["Method"] not in {"SGS_Lab", "Picketa_LENS"}:
            decision.update(disposition="held", reason="unknown_measurement_method")
            dispositions.append(decision)
            continue
        for name in selected:
            value = values[name]
            if name in formulas or value.strip() in {"", ".", "-9999", "NA"}:
                decision["null_cells"].append({"column": name, "raw": value if name not in formulas else None, "reason": "formula_excluded" if name in formulas else "source_missing_marker", "value": None})
                value = ""
            elif name in profile["measurements"] and safe_number(value) is None:
                decision["null_cells"].append({"column": name, "raw": value, "reason": "nonnumeric_measurement_held", "value": None})
                value = ""
            row[_alias(name)] = value
        row["evidence_kind"] = evidence
        decision["study_unit"] = group
        decision["prepared_row"] = len(groups[group]) + 2
        groups[group].append(row)
        dispositions.append(decision)
    excluded = []
    for i, name in enumerate(header, 1):
        if name not in selected:
            kind = "model_output" if "SoyMAP" in name else "derived" if name in {"RP", "YGD", "SGR", "YDH", "SN", "YDHR"} or profile["id"] == "onion-trial-yield" else "unselected"
            excluded.append({"column": name, "column_number": i, "kind": kind, "formula_cells": formula_count[name]})
    manifest = {"version": VERSION, "source_id": profile["id"], "raw_path": profile["path"], "raw_sha256": raw_sha,
        "dictionary": {"path": profile["dictionary"], "sha256": PINS[profile["dictionary"].split("#")[0]]}, "selected_sheet": profile["sheet"] or "CSV",
        "source_columns": [{"column": n, "column_number": header.index(n)+1, "output_column": _alias(n), "kind": "derived" if n in profile.get("derived", []) else "measurement" if n in profile["measurements"] else "context", "unit": profile["measurements"].get(n)} for n in selected],
        "excluded_columns": excluded, "group_columns": profile["group"], "formula_cells_excluded": dict(sorted(formula_count.items())),
        "notes": profile["notes"], "row_dispositions": dispositions, "disposition_counts": dict(Counter(d["disposition"] for d in dispositions)),
        "policies": {"formula_evaluation": False, "cached_formula_values": False, "geometry_inference": False, "farm_independence": "unknown", "runtime_rag_admission": False, "training_admission": False, "missing_output_encoding": "empty CSV cell, JSON null in disposition", "partial_dates": "unaltered text; not time-window eligible"}}
    return manifest, dict(sorted(groups.items()))


def prepare_sources(raw_root: Path, output_dir: Path) -> dict:
    """Verify pinned inputs and create a NEW review bundle; never modify originals."""
    raw_root, output_dir = Path(raw_root), Path(output_dir)
    if output_dir.exists():
        raise ValueError("output directory already exists; overwrites forbidden")
    inputs = {}
    for relative, expected in PINS.items():
        path = raw_root / relative
        if path.stat().st_size > MAX_SOURCE_BYTES:
            raise ValueError("source exceeds byte bounds")
        content = path.read_bytes()
        if _sha(content) != expected:
            raise ValueError(f"pinned SHA-256 mismatch: {relative}")
        inputs[relative] = content
    files = {}
    index = {"version": VERSION, "status": "prepared_for_review_only", "groups": [], "sources": [], "limitations": ["Study units are source partitions, not independent farms.", "No geometry, gold answers, runtime retrieval or training admission.", "Summary statistics are source-table diagnostics, not agronomic validation."]}
    for profile in PROFILES:
        manifest, groups = _prepare_profile(profile, inputs[profile["path"]])
        payload = _json(manifest)
        manifest_sha = _sha(payload)
        manifest_name = profile["id"] + ".manifest.json"
        files[manifest_name] = payload
        index["sources"].append({"source_id": profile["id"], "manifest": manifest_name, "manifest_sha256": manifest_sha, "raw_sha256": manifest["raw_sha256"], "disposition_counts": manifest["disposition_counts"], "group_count": len(groups)})
        for group, rows in groups.items():
            group_id = profile["id"] + "-" + _sha(group.encode())[:12]
            csv_name, mapping_name = group_id+".csv", group_id+".mapping.json"
            text = io.StringIO(newline="")
            writer = csv.DictWriter(text, fieldnames=list(rows[0]), lineterminator="\n")
            writer.writeheader()
            writer.writerows(rows)
            csv_bytes = text.getvalue().encode()
            parse_table(csv_name, csv_bytes)
            columns = [{"column": n, "label": n, "unit": None, "role": "identifier" if n in {"source_row", "study_unit"} else "context", "aggregation": "none", "evidence_role": "observation", "value_scope": "record"} for n in rows[0]]
            summaries = {}
            for c in columns:
                original = next((n for n in profile["measurements"] if _alias(n) == c["column"]), None)
                if original:
                    kind = "interpretation" if original in profile.get("derived", []) else rows[0]["evidence_kind"]
                    c.update(label=original, unit=profile["measurements"][original], role="measurement", aggregation="mean", evidence_role=kind)
                    numbers = [safe_number(r[c["column"]]) for r in rows]
                    present = [n for n in numbers if n is not None]
                    summaries[c["column"]] = {"count": len(present), "missing": len(numbers)-len(present), "mean": sum(present)/len(present) if present else None, "evidence_role": kind, "unit": c["unit"]}
            mapping = {"filters": {"study_unit": group}, "record_key": ["source_row"], "columns": columns,
                "source": {"title": f"{profile['id']}: {group}", "url": profile["url"], "license": profile["license"], "citation": f"{profile['notes']} Raw SHA256={manifest['raw_sha256']}; transformation manifest SHA256={manifest_sha}; {manifest_name}."}}
            files[csv_name], files[mapping_name] = csv_bytes, _json(mapping)
            index["groups"].append({"group_id": group_id, "source_id": profile["id"], "study_unit": group, "row_count": len(rows), "csv": csv_name, "csv_sha256": _sha(csv_bytes), "mapping": mapping_name, "mapping_sha256": _sha(files[mapping_name]), "raw_sha256": manifest["raw_sha256"], "manifest_sha256": manifest_sha, "summaries": summaries})
    files["index.json"] = _json(index)
    output_dir.mkdir(parents=True, exist_ok=False)
    for name, content in files.items():
        with (output_dir / name).open("xb") as stream:
            stream.write(content)
    return index
