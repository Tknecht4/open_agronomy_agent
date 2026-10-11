"""Bounded, offline CSCAP native-value preparation for the typed table importer.

No runtime RAG, training, gold or active-profile admission is performed here.
The source contract pins one public workbook and one site before extraction.
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import math
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any

from agronomy_agent.field_data import parse_table
from agronomy_agent.paths import repo_path

DEFAULT_CONTRACT = "data/manifests/cscap_arl_intake_v1.json"
MAX_RAW_BYTES = 13 * 1024 * 1024
MAX_EXPANDED_BYTES = 128 * 1024 * 1024
MAX_OUTPUT_BYTES = 512 * 1024
MAX_TABLE_ROWS = 5000
MAX_COLUMNS = 64


def _bytes(value: Any) -> bytes:
    return (json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(",", ":"), allow_nan=False) + "\n").encode()


def _sha(content: bytes) -> str:
    return hashlib.sha256(content).hexdigest()


def _pinned(root: Path, item: dict) -> bytes:
    name = item["name"]
    if Path(name).name != name:
        raise ValueError("source filename must be a basename")
    path = root / name
    if path.stat().st_size != item["bytes"] or path.stat().st_size > MAX_RAW_BYTES:
        raise ValueError(f"pinned size mismatch: {name}")
    content = path.read_bytes()
    if _sha(content) != item["sha256"]:
        raise ValueError(f"pinned SHA-256 mismatch: {name}")
    return content


def _literal(cell: Any) -> Any:
    if cell.data_type == "f":
        raise ValueError("selected source formulas cannot become observations")
    value = cell.value
    if isinstance(value, str) and len(value) > 32000:
        raise ValueError("source cell exceeds bounds")
    return value


def _table(book: Any, name: str, columns: list[str], definitions: dict | None = None) -> tuple[list[dict], dict[str, int]]:
    sheet = book[name]
    if sheet.max_row > MAX_TABLE_ROWS or sheet.max_column > MAX_COLUMNS:
        raise ValueError(f"source table exceeds bounds: {name}")
    iterator = sheet.iter_rows()
    header = [_literal(cell) for cell in next(iterator)]
    if len(header) != len(set(header)) or not set(columns) <= set(header):
        raise ValueError(f"missing or ambiguous source columns: {name}")
    positions = {column: header.index(column) for column in columns}
    description = next(iterator)
    units = next(iterator)
    if definitions:
        for column, spec in definitions.items():
            offset = positions[column]
            if _literal(description[offset]) != spec["description"] or _literal(units[offset]) != spec["unit_native"]:
                raise ValueError(f"source definition/unit mismatch: {column}")
    records = [{"row": row_number, "values": {column: _literal(cells[offset]) for column, offset in positions.items()}}
               for row_number, cells in enumerate(iterator, 4)
               if any(cell.value is not None and str(cell.value).strip() for cell in cells)]
    return records, {column: offset + 1 for column, offset in positions.items()}


def prepare_cscap_intake(raw_root: Path, output_dir: Path, *, contract_path: Path | None = None) -> dict:
    """Create a new small review bundle consumable by the existing typed importer."""
    output_dir = Path(output_dir)
    if output_dir.exists():
        raise FileExistsError("existing output directory; overwrites forbidden")
    contract_path = contract_path or repo_path(DEFAULT_CONTRACT)
    contract_content = Path(contract_path).read_bytes()
    contract = json.loads(contract_content)
    if contract.get("schema_version") != "open_agronomy_agent.cscap_intake.v1":
        raise ValueError("unsupported CSCAP source contract")
    if any(contract["admission"].get(key) is not False for key in ("runtime_rag", "active_profile", "training", "evaluation_gold")):
        raise ValueError("preparation cannot grant source admission")
    inputs = {role: _pinned(Path(raw_root), item) for role, item in contract["source_files"].items()}
    article = json.loads(inputs["article"])
    publisher_file = next((item for item in article.get("files", []) if item.get("id") == contract["source_files"]["workbook"]["publisher_file_id"]), {})
    if (article.get("id") != contract["article_id"] or article.get("version") != contract["article_version"]
            or article.get("doi") != contract["doi"] or article.get("license", {}).get("name") != "CC0"
            or article.get("license", {}).get("url") != contract["rights"]["licence_url"]
            or publisher_file.get("size") != len(inputs["workbook"])
            or publisher_file.get("is_link_only") is not False
            or publisher_file.get("supplied_md5") != contract["source_files"]["workbook"]["publisher_md5"]
            or hashlib.md5(inputs["workbook"]).hexdigest() != publisher_file.get("supplied_md5")):
        raise ValueError("publisher identity, file or CC0 scope mismatch")
    with zipfile.ZipFile(io.BytesIO(inputs["workbook"])) as archive:
        members = archive.infolist()
        if len(members) > 256 or sum(member.file_size for member in members) > MAX_EXPANDED_BYTES:
            raise ValueError("workbook archive exceeds bounds")
        if any(member.flag_bits & 1 or member.file_size > max(1, member.compress_size) * 100 for member in members):
            raise ValueError("workbook encryption or expansion exceeds bounds")
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(inputs["workbook"]), read_only=True, data_only=False, keep_links=False)
    selection, definitions = contract["selection"], contract["yield_columns"]
    site, years = selection["site"], selection["years"]
    try:
        plots, plot_columns = _table(book, "Plot Identifiers", ["uniqueid", "plotid", "rep", "tillage", "rotation", *[f"{year}crop" for year in years]])
        agronomic, agronomic_columns = _table(book, "Agronomic", ["uniqueid", "plotid", "year", *definitions], definitions)
    finally:
        book.close()
    plot_by_id = {}
    for record in plots:
        values = record["values"]
        if values["uniqueid"] != site:
            continue
        key = str(values["plotid"] or "")
        if not key or key in plot_by_id:
            raise ValueError("blank or duplicate source plot identity")
        plot_by_id[key] = record
    if len(plot_by_id) != selection["expected_plot_count"]:
        raise ValueError("selected plot count mismatch")
    raw_sha = _sha(inputs["workbook"])
    rows, dispositions, seen = [], [], set()
    for record in agronomic:
        values = record["values"]
        if values["uniqueid"] != site:
            continue
        year, plot_id = values["year"], str(values["plotid"] or "")
        key = (plot_id, year)
        if type(year) is not int or year not in years or key in seen or plot_id not in plot_by_id:
            raise ValueError("unknown, duplicate or unjoinable source plot-year")
        seen.add(key)
        plot = plot_by_id[plot_id]
        crop = plot["values"][f"{year}crop"]
        row = {"study_site": site, "plot_id": plot_id, "year": str(year), "replicate": str(plot["values"]["rep"]),
               "native_crop": crop, "tillage_code": plot["values"]["tillage"], "rotation_code": plot["values"]["rotation"]}
        disposition = {"agronomic_row": record["row"], "plot_identifier_row": plot["row"], "null_cells": []}
        for column, spec in definitions.items():
            value = values[column]
            if value is None or isinstance(value, str) and not value.strip():
                row[spec["output_column"]] = ""
                disposition["null_cells"].append({"column": column, "raw": value, "value": None})
            elif type(value) not in (int, float) or not math.isfinite(value) or value < 0:
                raise ValueError(f"unreviewed nonnumeric or invalid yield: {column}")
            elif crop != spec["crop"]:
                raise ValueError("source crop/yield column conflict")
            else:
                row[spec["output_column"]] = str(value)
        row["source_row"] = f"sha256:{raw_sha}#Agronomic!row={record['row']}"
        row["plot_source_row"] = f"sha256:{raw_sha}#Plot Identifiers!row={plot['row']}"
        rows.append(row)
        dispositions.append(disposition)
    year_counts = Counter(row["year"] for row in rows)
    if len(rows) != selection["expected_record_count"] or year_counts != Counter({str(year): selection["expected_records_per_year"] for year in years}):
        raise ValueError("selected plot-year counts mismatch")
    rows.sort(key=lambda row: (row["year"], row["plot_id"]))
    text = io.StringIO(newline="")
    writer = csv.DictWriter(text, fieldnames=list(rows[0]), lineterminator="\n")
    writer.writeheader()
    writer.writerows(rows)
    csv_content = text.getvalue().encode()
    parse_table("cscap-arl.csv", csv_content)
    columns = [{"column": name, "label": name, "unit": None, "role": "identifier" if name in {"source_row", "plot_source_row", "plot_id"} else "context",
                "aggregation": "none", "evidence_role": "observation", "value_scope": "record"} for name in rows[0]]
    for column in columns:
        spec = next((spec for spec in definitions.values() if spec["output_column"] == column["column"]), None)
        if spec:
            column.update(label=spec["description"], unit=spec["unit"], role="measurement", aggregation="mean")
    receipt = {"schema_version": "open_agronomy_agent.cscap_prepared_intake.v1", "status": "prepared_for_review_only",
               "extractor": {"name": "agronomy_agent.cscap_intake", "version": 1, "implementation_sha256": _sha(Path(__file__).read_bytes())},
               "source_id": contract["source_id"], "source_contract_sha256": _sha(contract_content), "raw_sha256": raw_sha,
               "article_sha256": _sha(inputs["article"]), "source_files_copied": False, "row_count": len(rows), "plot_count": len(plot_by_id),
               "year_counts": dict(sorted(year_counts.items())), "numeric_counts": {spec["output_column"]: sum(bool(row[spec["output_column"]]) for row in rows) for spec in definitions.values()},
               "definitions": definitions, "source_columns": {"Agronomic": agronomic_columns, "Plot Identifiers": plot_columns},
               "definition_rows": {"descriptions": 2, "units": 3}, "row_dispositions": dispositions, "admission": contract["admission"], "limits": contract["limits"],
               "formula_evaluation": False, "imputation": False, "unit_conversion": False, "runtime_rag_chunks": 0,
               "groups": [], "files": {"cscap-arl.csv": {"bytes": len(csv_content), "sha256": _sha(csv_content)}}}
    files = {"cscap-arl.csv": csv_content}
    # Keep the product's explicit plot/field identity guard. One CSV is shared;
    # each mapping admits exactly one experimental plot and its dated records.
    from agronomy_agent.server.storage.field_data_store import _validate_mapping
    parsed = parse_table("cscap-arl.csv", csv_content)
    for plot_id in sorted(plot_by_id):
        name = f"cscap-arl-{plot_id}.mapping.json"
        if Path(name).name != name:
            raise ValueError("unsafe source plot identifier")
        mapping = {"filters": {"study_site": site, "plot_id": plot_id}, "record_key": ["study_site", "plot_id", "year"], "columns": columns,
                   "source": {"title": f"{contract['title']} :: source plot {plot_id}", "url": contract["url"], "license": contract["rights"]["identifier"],
                              "citation": f"{contract['rights']['attribution']} Raw SHA256={raw_sha}; source contract SHA256={_sha(contract_content)}. Historical repeated study plots; unweighted row means are not treatment effects or current farmer-field estimates. No geometry or field-operation joins."}}
        _, selected = _validate_mapping(mapping, parsed)
        files[name] = _bytes(mapping)
        receipt["files"][name] = {"bytes": len(files[name]), "sha256": _sha(files[name])}
        receipt["groups"].append({"study_site": site, "plot_id": plot_id, "row_count": len(selected), "csv": "cscap-arl.csv", "mapping": name})
    files["source-receipt.json"] = _bytes(receipt)
    if sum(map(len, files.values())) > MAX_OUTPUT_BYTES:
        raise ValueError("prepared output exceeds byte budget")
    output_dir.mkdir(parents=True, exist_ok=False)
    for name, content in files.items():
        with (output_dir / name).open("xb") as stream:
            stream.write(content)
    return receipt
