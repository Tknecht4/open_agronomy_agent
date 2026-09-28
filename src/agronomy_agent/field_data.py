"""Bounded, lossless-enough table intake for private field uploads.

Parsed values are strings. A caller must explicitly interpret columns, units and
field scope before any row can become queryable.
"""

from __future__ import annotations

import csv
import hashlib
import io
import math
import re
import zipfile
from pathlib import PurePath
from typing import Any
from xml.etree.ElementTree import ParseError

MAX_SOURCE_BYTES = 8 * 1024 * 1024
MAX_ROWS = 50_000
MAX_COLUMNS = 128
MAX_CELL_CHARS = 32_000
MAX_ZIP_MEMBERS = 256
MAX_ZIP_EXPANDED_BYTES = 64 * 1024 * 1024
FORMULA_PREFIXES = ("=", "+", "-", "@")


def _safe_filename(filename: str) -> str:
    name = PurePath(filename.replace("\\", "/")).name
    if not name or name in {".", ".."} or len(name) > 255:
        raise ValueError("invalid filename")
    return name


def _check_header(header: list[str]) -> None:
    if not header or len(header) > MAX_COLUMNS:
        raise ValueError("header is empty or exceeds column limit")
    if any(not name.strip() or len(name) > 256 for name in header):
        raise ValueError("header contains blank or overlong names")
    if len(set(header)) != len(header):
        raise ValueError("duplicate column names")


def _check_row(row: list[str], width: int) -> None:
    if len(row) != width:
        raise ValueError("ragged row")
    if any(len(value) > MAX_CELL_CHARS for value in row):
        raise ValueError("cell exceeds size limit")
    # Spreadsheet formula syntax is never evaluated or silently rendered.
    if any(value.lstrip().startswith(FORMULA_PREFIXES) and safe_number(value) is None for value in row):
        raise ValueError("formula-like cell requires a source-safe replacement")


def _text_table(content: bytes, *, filename: str, encoding: str | None) -> tuple[list[str], list[dict[str, Any]], str, str]:
    if encoding is None:
        encoding = "utf-8-sig"
    if encoding.lower().replace("_", "-") not in {"utf-8", "utf-8-sig", "latin-1", "iso-8859-1"}:
        raise ValueError("unsupported encoding; select utf-8 or latin-1 explicitly")
    try:
        text = content.decode(encoding, errors="strict")
    except UnicodeError as exc:
        raise ValueError("text decoding failed; select the source encoding explicitly") from exc
    if "\x00" in text:
        raise ValueError("NUL byte in text table")
    sample = text[: min(16_384, len(text))]
    expected = "\t" if filename.lower().endswith((".tsv", ".tab")) else None
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",\t;|")
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = expected or ","
    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter, strict=True)
    try:
        header = next(reader)
    except (StopIteration, csv.Error) as exc:
        raise ValueError("table has no valid header") from exc
    _check_header(header)
    rows: list[dict[str, Any]] = []
    try:
        for record_number, values in enumerate(reader, start=2):
            if record_number > MAX_ROWS + 1:
                raise ValueError("table exceeds row limit")
            _check_row(values, len(header))
            rows.append({"locator": {"record": record_number, "physical_line_end": reader.line_num}, "values": dict(zip(header, values, strict=True))})
    except csv.Error as exc:
        raise ValueError("malformed delimited table") from exc
    return header, rows, encoding, delimiter


def _xlsx_table(content: bytes) -> tuple[list[str], list[dict[str, Any]], str, str]:
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS or sum(m.file_size for m in members) > MAX_ZIP_EXPANDED_BYTES:
                raise ValueError("workbook archive exceeds limits")
            names = set(archive.namelist())
            if not {"[Content_Types].xml", "xl/workbook.xml"} <= names:
                raise ValueError("invalid XLSX workbook structure")
            for member in members:
                if member.flag_bits & 1:
                    raise ValueError("encrypted XLSX members are unsupported")
                if member.file_size > MAX_ZIP_EXPANDED_BYTES or (member.compress_size and member.file_size > 200 * member.compress_size):
                    raise ValueError("workbook member exceeds expansion limit")
    except (zipfile.BadZipFile, zipfile.LargeZipFile) as exc:
        raise ValueError("invalid XLSX archive") from exc
    try:
        import openpyxl
    except ImportError as exc:
        raise ValueError("XLSX support requires openpyxl") from exc
    malformed_workbook_errors = (KeyError, zipfile.BadZipFile, openpyxl.utils.exceptions.InvalidFileException,
                                 ParseError, EOFError, UnicodeError)
    try:
        workbook = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except malformed_workbook_errors as exc:
        raise ValueError("invalid XLSX workbook structure") from exc
    try:
        if len(workbook.sheetnames) != 1:
            raise ValueError("XLSX must contain exactly one sheet in this pilot")
        sheet = workbook.active
        iterator = sheet.iter_rows()
        first = next(iterator, None)
        if first is None:
            raise ValueError("table has no header")
        header = ["" if cell.value is None else str(cell.value) for cell in first]
        _check_header(header)
        rows: list[dict[str, Any]] = []
        for record_number, cells in enumerate(iterator, start=2):
            if record_number > MAX_ROWS + 1:
                raise ValueError("table exceeds row limit")
            if any(cell.data_type == "f" for cell in cells):
                raise ValueError("XLSX formula cell has ambiguous cached value")
            values = ["" if cell.value is None else str(cell.value) for cell in cells]
            # openpyxl pads rows to sheet width; trim only trailing empty cells.
            if len(values) > len(header) and any(values[len(header):]):
                raise ValueError("ragged row")
            values = values[:len(header)] + [""] * max(0, len(header) - len(values))
            _check_row(values, len(header))
            rows.append({"locator": {"sheet": sheet.title, "record": record_number}, "values": dict(zip(header, values, strict=True))})
        return header, rows, "xlsx", ""
    except malformed_workbook_errors as exc:
        raise ValueError("invalid XLSX workbook structure") from exc
    finally:
        workbook.close()


def parse_table(filename: str, content: bytes, encoding: str | None = None) -> dict[str, Any]:
    """Parse one bounded source file; preserve source strings and record locators."""
    filename = _safe_filename(filename)
    if not content or len(content) > MAX_SOURCE_BYTES:
        raise ValueError("source is empty or exceeds byte limit")
    suffix = filename.lower().rsplit(".", 1)[-1]
    if suffix == "xlsx":
        if encoding is not None:
            raise ValueError("XLSX does not take a text encoding")
        columns, rows, selected_encoding, delimiter = _xlsx_table(content)
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                worksheet_members = [name for name in archive.namelist() if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", name)]
                if len(worksheet_members) != 1:
                    raise ValueError("XLSX must contain exactly one worksheet member")
                member_name = worksheet_members[0]
                member_sha256 = hashlib.sha256(archive.read(member_name)).hexdigest()
        except zipfile.BadZipFile as exc:
            raise ValueError("invalid XLSX workbook structure") from exc
    elif suffix in {"csv", "tsv", "tab"}:
        columns, rows, selected_encoding, delimiter = _text_table(content, filename=filename, encoding=encoding)
        member_name = filename
        member_sha256 = hashlib.sha256(content).hexdigest()
    else:
        raise ValueError("unsupported table format")
    warnings = []
    if suffix in {"tab", "tsv"} and delimiter != "\t":
        warnings.append("Filename suggests tabs but content uses another delimiter; parsed delimiter is recorded.")
    if suffix == "csv" and delimiter != ",":
        warnings.append("Filename suggests commas but content uses another delimiter; parsed delimiter is recorded.")
    return {"filename": filename, "member_name": member_name, "member_sha256": member_sha256,
            "columns": columns, "rows": rows, "row_count": len(rows), "encoding": selected_encoding,
            "delimiter": delimiter, "warnings": warnings}


def safe_number(value: str) -> float | None:
    """Return a finite unambiguous decimal or None; caller tracks missing/invalid."""
    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?", value.strip()):
        return None
    try:
        number = float(value)
    except (ValueError, OverflowError):
        return None
    return number if math.isfinite(number) else None
